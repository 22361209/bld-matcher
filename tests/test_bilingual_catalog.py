from __future__ import annotations

import io
import json
import sqlite3
import tarfile
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from werkzeug.datastructures import MultiDict

from tests.web_app_test_base import WebAppTestBase


class BilingualCatalogTest(WebAppTestBase):
    def setUp(self):
        self.login()
        self.addCleanup(self.cleanup_products, "CATLANG-%")
        from app.modules.products.persistence import upsert_product

        with self.web.connect(self.web.DB_PATH) as connection:
            for number, (english, chinese) in enumerate(((1, 1), (1, 0), (0, 1), (0, 0))):
                upsert_product(connection, {
                    "bld_no": f"CATLANG-{number}", "series": "CATLANG BRAND",
                    "item": "Front Lower Control Arm", "item_zh": "前下控制臂",
                    "models": "Foreign Model 2010–2015", "models_zh": "中文车型 2018–2023" if number == 0 else "",
                    "oe_no_1": f"LANG-OE-{number}", "product_status": "1 个球头",
                    "catalog_visible_en": english, "catalog_visible_zh": chinese,
                }, actor="test")
            connection.commit()

    def _records(self):
        with self.web.connect(self.web.DB_PATH) as connection:
            return {row["bld_no"]: dict(row) for row in connection.execute(
                "SELECT * FROM products WHERE bld_no LIKE 'CATLANG-%'"
            )}

    def test_language_visibility_and_management_preserve_all_four_combinations(self):
        for language, expected in (("en", {0, 1}), ("zh", {0, 2})):
            for endpoint in ("/products", "/products/fragment"):
                response = self.client.get(endpoint, query_string={"bld": "CATLANG-", "lang": language})
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                for number in range(4):
                    self.assertEqual(f"CATLANG-{number}" in html, number in expected)
                self.assertIn(f"lang={language}", html)
                self.assertIn("data-products-mobile-cards", html)
        all_rows = self.client.get("/products", query_string={"bld": "CATLANG-", "visibility": "all"}).text
        self.assertTrue(all(f"CATLANG-{number}" in all_rows for number in range(4)))
        hidden = self.client.get("/products", query_string={"bld": "CATLANG-", "visibility": "neither"}).text
        self.assertIn("CATLANG-3", hidden)
        self.assertNotIn("CATLANG-0", hidden)

    def test_chinese_visible_cells_never_copy_english_models(self):
        response = self.client.get("/products/fragment", query_string={"bld": "CATLANG-", "lang": "zh"})
        html = response.text
        self.assertIn('data-col="models">中文车型 2018–2023', html)
        self.assertIn('data-col="models">未填写', html)
        self.assertRegex(html, r'data-col="item">\s*<span>前下控制臂')
        self.assertIn('product-mobile-card-item" title="前下控制臂"', html)
        self.assertIn('product-mobile-card-model" title="未填写">车型 · 未填写', html)
        # Original values remain in copy payloads; visible model cells are independent.
        self.assertNotIn('data-col="models">Foreign Model', html)
        self.assertNotIn('product-mobile-card-model" title="Foreign Model', html)

    def test_save_round_trip_unchecked_flags_and_legacy_save(self):
        data = MultiDict([
            ("bld_no", "CATLANG-0"), ("series", "CATLANG BRAND"),
            ("item", "Front Lower Control Arm"), ("item_zh", "自定中文名"),
            ("models", "Foreign 2000–2008"), ("models_zh", "国产车型 2012–2016"),
            ("active", "1"), ("catalog_visible_en", "0"),
            ("catalog_visible_zh", "0"), ("catalog_visible_zh", "1"),
        ])
        response = self.client.post("/products/save", data=data)
        self.assertEqual(response.status_code, 302)
        saved = self._records()["CATLANG-0"]
        self.assertEqual((saved["catalog_visible_en"], saved["catalog_visible_zh"]), (0, 1))
        self.assertEqual((saved["item_zh"], saved["models_zh"]), ("自定中文名", "国产车型 2012–2016"))
        self.client.post("/products/save", data={
            "bld_no": "CATLANG-0", "item": "Front Lower Control Arm",
            "models": "Foreign updated", "active": "1",
        })
        retained = self._records()["CATLANG-0"]
        for field in ("item_zh", "models_zh", "catalog_visible_en", "catalog_visible_zh"):
            self.assertEqual(retained[field], saved[field])
        form = self.client.get(f"/products/{saved['id']}/edit").text
        self.assertLess(form.index("产品名称（英文）"), form.index("产品名称（中文）"))
        self.assertIn("国产车型 2012–2016", form)

    def test_new_product_defaults_and_explicit_empty_chinese_models(self):
        self.client.post("/products/save", data={
            "bld_no": "CATLANG-NEW", "item": "Front Lower Control Arm", "models": "English only", "active": "1",
        })
        row = self._records()["CATLANG-NEW"]
        self.assertEqual((row["catalog_visible_en"], row["catalog_visible_zh"]), (1, 1))
        self.assertEqual(row["item_zh"], "前下控制臂")
        self.assertEqual(row["models_zh"], "")
        new_html = self.client.get("/products/new").text
        self.assertIn('name="catalog_visible_en" value="1" checked', new_html)
        self.assertIn('name="catalog_visible_zh" value="1" checked', new_html)

    def test_both_export_formats_choose_language_independently_and_report_blanks(self):
        for export_format in ("bld", "brand"):
            for language, expected in (("en", {"CATLANG-0", "CATLANG-1"}), ("zh", {"CATLANG-0", "CATLANG-2"})):
                response = self.client.post("/products/export", data={
                    "bld": "CATLANG-", "lang": "en", "export_language": language,
                    "visibility": "visible", "status": "active", "export_format": export_format,
                })
                self.assertEqual(response.status_code, 200)
                self.addCleanup(response.close)
                workbook = load_workbook(io.BytesIO(response.data))
                self.addCleanup(workbook.close)
                rows = list(workbook.active.values)
                headers = rows[0]
                bld_column = headers.index("BLD NO.")
                self.assertEqual({row[bld_column] for row in rows[1:]}, expected)
                if language == "zh":
                    self.assertEqual(workbook["目录说明"]["B4"].value, 1)
                    model_column = headers.index("中文车型")
                    self.assertEqual({row[model_column] for row in rows[1:]}, {"中文车型 2018–2023", None})
                    name_column = headers.index("中文产品名称")
                    self.assertTrue(all(row[name_column] == "前下控制臂" for row in rows[1:]))
                else:
                    self.assertTrue(all(row[headers.index("Models")] == "Foreign Model 2010–2015" for row in rows[1:]))
                    self.assertIn("Product Status", headers)
                self.assertEqual(workbook["目录说明"]["B2"].value, language)

    def test_invalid_catalog_language_is_rejected_before_render_or_export(self):
        for endpoint in ("/products", "/products/fragment", "/products/export"):
            self.assertEqual(self.client.get(endpoint, query_string={"lang": "invalid"}).status_code, 400)
        self.assertEqual(self.client.post("/products/export", data={"export_language": "invalid"}).status_code, 400)

    def test_chinese_display_export_cannot_overwrite_english_through_either_importer(self):
        from app.matcher import ProductCatalog
        from app.modules.products.catalog_import import CatalogImportChoices, read_catalog_import

        response = self.client.post("/products/export", data={"bld": "CATLANG-", "export_language": "zh"})
        self.addCleanup(response.close)
        path = self.root / "chinese-display.xlsx"
        path.write_bytes(response.data)
        for importer in (
            lambda: ProductCatalog.from_excel(path),
            lambda: read_catalog_import(path, choices=CatalogImportChoices(("CATLANG BRAND",), ("Front Lower Control Arm",))),
        ):
            with self.assertRaisesRegex(ValueError, "中文展示目录"):
                importer()
        self.assertEqual(self._records()["CATLANG-0"]["models"], "Foreign Model 2010–2015")

    def test_catalog_source_and_optional_import_columns_round_trip(self):
        from app.modules.products.catalog_import import CatalogImportChoices, read_catalog_import
        from app.modules.products.catalog_template import build_catalog_import_template
        from app.modules.products.repository import SQLiteProductRepository
        from app.modules.products.persistence import import_catalog

        source = self.root / "catalog-source.xlsx"
        target = self.root / "restored.sqlite3"
        with self.web.connect(self.web.DB_PATH) as connection:
            SQLiteProductRepository(connection, self.web.DB_PATH).export_catalog_source(source)
        with self.web.connect(target) as connection:
            import_catalog(connection, source)
            restored = connection.execute("SELECT * FROM products WHERE bld_no='CATLANG-3'").fetchone()
            self.assertEqual((restored["catalog_visible_en"], restored["catalog_visible_zh"]), (0, 0))
            restored = connection.execute("SELECT * FROM products WHERE bld_no='CATLANG-0'").fetchone()
            self.assertEqual(restored["models_zh"], "中文车型 2018–2023")
        choices = CatalogImportChoices(("CATLANG BRAND",), ("Front Lower Control Arm",))
        workbook = load_workbook(build_catalog_import_template(choices))
        headers = [cell.value for cell in workbook.active[1]]
        values = {"BLD NO.": "CATLANG-IMPORT", "SERIES": "CATLANG BRAND", "ITEM": "Front Lower Control Arm",
                  "OE NO.1": "OE", "Models": "English", "产品状态": "1球头", "导入单价": 1,
                  "中文产品名称": "导入中文", "中文车型": "不同年份 2020–2025",
                  "英文目录可见": 0, "中文目录可见": 1}
        workbook.active.append([values.get(header) for header in headers])
        path = self.root / "optional-import.xlsx"
        workbook.save(path)
        workbook.close()
        imported = read_catalog_import(path, choices=choices)[0]
        self.assertEqual(imported.values()["models_zh"], "不同年份 2020–2025")
        self.assertEqual(imported.values()["catalog_visible_en"], 0)

    def test_legacy_excel_import_preserves_independent_catalog_fields(self):
        from app.modules.products.catalog_import import CatalogImportChoices, read_catalog_import, build_catalog_import_preview
        from app.modules.products.factory import get_product_service
        from openpyxl import Workbook

        workbook = Workbook()
        workbook.active.append(["BLD NO.", "SERIES", "ITEM", "OE NO.1", "Models", "产品状态", "导入单价"])
        workbook.active.append(["CATLANG-3", "CATLANG BRAND", "Front Lower Control Arm", "NEW-OE", "Changed English", "1球头", 1])
        path = self.root / "legacy.xlsx"
        workbook.save(path)
        workbook.close()
        rows = read_catalog_import(path, choices=CatalogImportChoices(("CATLANG BRAND",), ("Front Lower Control Arm",)))
        existing = get_product_service().find_by_bld("CATLANG-3")
        preview = build_catalog_import_preview(rows, {"CATLANG-3": existing})
        labels = {field["label"] for field in preview.conflicts[0].fields}
        self.assertNotIn("中文车型", labels)
        get_product_service().save(rows[0].values(), actor="test")
        row = self._records()["CATLANG-3"]
        self.assertEqual((row["catalog_visible_en"], row["catalog_visible_zh"]), (0, 0))
        self.assertEqual(row["item_zh"], "前下控制臂")

    def test_filters_keep_stable_english_keys_and_search_chinese(self):
        response = self.client.get("/products/fragment", query_string={
            "bld": "CATLANG-", "lang": "zh", "item": "Front Lower Control Arm",
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn('value="Front Lower Control Arm"', response.text)
        self.assertIn("前下控制臂", response.text)
        search = self.client.get("/products/fragment", query_string={"bld": "中文车型", "lang": "zh"})
        self.assertIn("CATLANG-0", search.text)
        self.assertNotIn("CATLANG-2", search.text)


class BilingualCatalogMigrationTest(unittest.TestCase):
    def test_legacy_rows_default_both_visible_and_do_not_translate_models(self):
        from app.migrations import _add_bilingual_catalog

        connection = sqlite3.connect(":memory:")
        connection.row_factory = sqlite3.Row
        self.addCleanup(connection.close)
        connection.execute("CREATE TABLE products (item TEXT, models TEXT)")
        connection.executemany("INSERT INTO products VALUES (?, ?)", [
            ("Front Lower Control Arm", "Foreign 2005–2010"), ("Central Link", "Other")])
        _add_bilingual_catalog(connection)
        rows = connection.execute(
            "SELECT item_zh, models_zh, catalog_visible_en, catalog_visible_zh FROM products"
        ).fetchall()
        self.assertEqual([tuple(row) for row in rows], [("前下控制臂", "", 1, 1), ("", "", 1, 1)])
        connection.execute("UPDATE products SET item_zh='用户译名', models_zh='用户车型', catalog_visible_en=0")
        _add_bilingual_catalog(connection)
        self.assertEqual(tuple(connection.execute(
            "SELECT item_zh, models_zh, catalog_visible_en FROM products LIMIT 1"
        ).fetchone()), ("用户译名", "用户车型", 0))


class BilingualBusinessSyncTest(unittest.TestCase):
    def test_new_package_round_trip_and_missing_legacy_fields_preserve_local_values(self):
        from app.database import connect
        from app.modules.products.persistence import upsert_product
        from app.modules.business_sync.infrastructure import BusinessSyncRepository
        from app.modules.business_sync.service import BusinessSyncService

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source, target = root / "source.sqlite3", root / "target.sqlite3"
            with connect(source) as connection:
                upsert_product(connection, {"bld_no": "LANGSYNC", "item": "Front Lower Control Arm",
                    "item_zh": "已核对中文", "models_zh": "独立车型 2025", "catalog_visible_en": 0}, actor="test")
                connection.commit()
            service = BusinessSyncService(BusinessSyncRepository(source))
            package = service.export(output_path=root / "new.tar.gz", selected=("products",), actor="test")
            receiving = BusinessSyncService(BusinessSyncRepository(target))
            preview = receiving.preview(package)
            receiving.apply(package, backup_path=root / "before.sqlite3", actor="test", expected_token=preview["token"])
            with connect(target) as connection:
                row = dict(connection.execute("SELECT * FROM products WHERE bld_no='LANGSYNC'").fetchone())
                self.assertEqual((row["item_zh"], row["models_zh"], row["catalog_visible_en"]),
                                 ("已核对中文", "独立车型 2025", 0))
            with tarfile.open(package) as archive:
                payload = json.load(archive.extractfile("data.json"))
            for field in ("item_zh", "models_zh", "catalog_visible_en", "catalog_visible_zh"):
                payload["products"][0].pop(field)
            payload["products"][0]["models"] = "Legacy changed English"
            payload["products"][0]["updated_at"] = "2099-01-01 00:00:00"
            legacy = root / "legacy.tar.gz"
            with tarfile.open(legacy, "w:gz") as archive:
                manifest = {"package_type": "bld_business_data", "version": 1, "datasets": ["products"]}
                for name, document in (("manifest.json", manifest), ("data.json", payload)):
                    data = json.dumps(document).encode()
                    member = tarfile.TarInfo(name)
                    member.size = len(data)
                    archive.addfile(member, io.BytesIO(data))
            preview = receiving.preview(legacy)
            receiving.apply(legacy, backup_path=root / "legacy-before.sqlite3", actor="test", expected_token=preview["token"])
            with connect(target) as connection:
                row = connection.execute("SELECT * FROM products WHERE bld_no='LANGSYNC'").fetchone()
                self.assertEqual((row["item_zh"], row["models_zh"], row["catalog_visible_en"]),
                                 ("已核对中文", "独立车型 2025", 0))
                self.assertEqual(row["models"], "Legacy changed English")
